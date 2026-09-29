"""A synthetic corpus shaped like the one the Mini-Challenge 3 spec describes, plus a scripted model.

Built from the published description (file types, the four hostile cases, the ten sample
questions), not from the organisers' starter kit. Everything is generated at test time so no
binary fixtures are checked in.
"""

from __future__ import annotations

import hashlib
import io
import os
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

__all__ = [
    "QUESTIONS",
    "ScriptedModel",
    "build_corpus",
    "make_pdf",
]

# (question, expected answer, expected citations)
QUESTIONS: list[tuple[str, str, list[str]]] = [
    (
        "What is the maximum junction temperature of the TQ-40?",
        "94",
        ["specs/tq40_datasheet_r2.pdf"],
    ),
    (
        "In which quarter does the TQ-60 enter customer sampling?",
        "Q3 FY27",
        ["planning/roadmap_fy27.docx"],
    ),
    (
        "What is the part number of the field-replaceable fan assembly for the TQ-40?",
        "ORR-FAN-2214-B",
        ["support/rma_parts.xlsx"],
    ),
    ("Which firmware version fixed ticket ORR-1847?", "4.3.2", ["support/bug_database.csv"]),
    (
        "What error code is logged when the thermal throttle engages?",
        "E7731",
        ["logs/prod_inference_2026-09-02.log"],
    ),
    (
        "What is the default batch timeout, in seconds, in the ingest service?",
        "180",
        ["engineering/ingest_service.py"],
    ),
    (
        "Which backplane pin carries THERM_ALERT# on the TQ-40?",
        "B14",
        ["specs/backplane_pinout.png"],
    ),
    (
        "What board revision is printed on the asset label?",
        "REV-C2",
        ["support/asset_label.jpg"],
    ),
    (
        "The production log shows a thermal throttle incident. Which firmware release fixed the "
        "underlying defect?",
        "4.3.2",
        ["logs/prod_inference_2026-09-02.log", "support/bug_database.csv"],
    ),
    ("What is the unit price of the TQ-40 at 10,000 unit volume?", "", []),
]

IMAGE_TRANSCRIPTS = {
    "specs/backplane_pinout.png": "TQ-40 BACKPLANE PINOUT\nPin B12: PWR_GOOD\nPin B14: THERM_ALERT#\nPin B15: FAN_TACH",
    "support/asset_label.jpg": "ASSET LABEL\nProduct: TQ-40 Compute Blade\nBoard Rev: REV-C2\nS/N: 000317",
}


def _escape(text: str) -> str:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def make_pdf(pages: list[list[str]]) -> bytes:
    """A minimal, valid, text-bearing PDF (ASCII only) that pypdf can extract from."""
    count = len(pages)
    font_id = 3 + 2 * count
    kids = " ".join(f"{3 + 2 * i} 0 R" for i in range(count))
    objects: dict[int, str] = {
        1: "<< /Type /Catalog /Pages 2 0 R >>",
        2: f"<< /Type /Pages /Kids [{kids}] /Count {count} >>",
        font_id: "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    }
    for i, lines in enumerate(pages):
        page_id, content_id = 3 + 2 * i, 4 + 2 * i
        body = (
            "BT /F1 12 Tf 50 740 Td 16 TL "
            + " ".join(f"({_escape(ln)}) Tj T*" for ln in lines)
            + " ET"
        )
        objects[page_id] = (
            "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Contents {content_id} 0 R /Resources << /Font << /F1 {font_id} 0 R >> >> >>"
        )
        objects[content_id] = f"<< /Length {len(body)} >>\nstream\n{body}\nendstream"
    out = bytearray(b"%PDF-1.4\n")
    offsets: dict[int, int] = {}
    for number in sorted(objects):
        offsets[number] = len(out)
        out += f"{number} 0 obj\n{objects[number]}\nendobj\n".encode("latin-1")
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for number in sorted(objects):
        out += f"{offsets[number]:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return bytes(out)


def _encrypt_pdf(raw: bytes, password: str) -> bytes:
    from pypdf import PdfReader, PdfWriter

    writer = PdfWriter(clone_from=PdfReader(io.BytesIO(raw)))
    writer.encrypt(user_password=password, algorithm="RC4-128")
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def _image(path: Path, text: str, fmt: str) -> None:
    from PIL import Image, ImageDraw

    image = Image.new("RGB", (640, 220), "white")
    draw = ImageDraw.Draw(image)
    for row, line in enumerate(text.split("\n")):
        draw.text((12, 12 + row * 24), line, fill="black")
    image.save(path, format=fmt)


def build_corpus(root: Path, *, include_unreadable: bool = False) -> Path:
    """Create the sample corpus under ``root`` and return it."""
    from docx import Document
    from openpyxl import Workbook

    def put(relative: str) -> Path:
        target = root / relative
        target.parent.mkdir(parents=True, exist_ok=True)
        return target

    put("archive").mkdir(exist_ok=True)  # the empty directory

    put("specs/tq40_datasheet_r2.pdf").write_bytes(
        make_pdf(
            [
                [
                    "TQ-40 Datasheet, Revision 2",
                    "Absolute maximum ratings",
                    "Supply voltage: 12 V nominal",
                ],
                [
                    "Thermal characteristics",
                    "Maximum junction temperature: 94 deg C",
                    "Operating range 0 to 70 deg C",
                ],
            ]
        )
    )
    put("specs/tq40_datasheet_r1_WITHDRAWN.pdf").write_bytes(
        make_pdf(
            [
                ["WITHDRAWN - superseded by Revision 2", "TQ-40 Datasheet, Revision 1"],
                ["Thermal characteristics", "Maximum junction temperature: 85 deg C"],
            ]
        )
    )

    document = Document()
    document.add_heading("FY27 Roadmap", level=1)
    document.add_paragraph("Planning assumptions are reviewed every quarter.")
    table = document.add_table(rows=1, cols=3)
    for cell, title in zip(
        table.rows[0].cells, ["Milestone", "Product", "Quarter"], strict=True
    ):
        cell.text = title
    for row in [
        ("Customer sampling", "TQ-60", "Q3 FY27"),
        ("Volume production", "TQ-60", "Q1 FY28"),
        ("Customer sampling", "TQ-40", "Q1 FY26"),
    ]:
        cells = table.add_row().cells
        for cell, value in zip(cells, row, strict=True):
            cell.text = value
    document.save(str(put("planning/roadmap_fy27.docx")))

    workbook = Workbook()
    parts = workbook.active
    assert parts is not None
    parts.title = "Parts"
    parts.append(["part_number", "description", "product", "lead_time_days"])
    parts.append(["ORR-FAN-2214-B", "Field-replaceable fan assembly", "TQ-40", 14])
    parts.append(["ORR-PSU-1100-A", "Power supply module", "TQ-40", 21])
    costs = workbook.create_sheet("Costs")
    costs.append(["part_number", "unit_cost_usd"])
    costs.append(["ORR-FAN-2214-B", 38.5])
    workbook.save(str(put("support/rma_parts.xlsx")))

    put("support/bug_database.csv").write_text(
        "ticket,title,severity,status,fixed_in\n"
        "ORR-1847,Thermal throttle engages below rated load,high,closed,4.3.2\n"
        "ORR-1850,Fan tachometer reads zero after hot swap,medium,closed,4.4.0\n"
        "ORR-1862,Log rotation stalls on full disk,low,open,\n",
        encoding="utf-8",
    )
    put("logs/prod_inference_2026-09-02.log").write_text(
        "\n".join(
            [
                f"2026-09-02 03:{m:02d}:00 INFO batch {m} complete latency_ms={40 + m}"
                for m in range(10, 21)
            ]
            + [
                "2026-09-02 03:21:44 WARN thermal throttle engaged error_code=E7731 (see ORR-1847)"
            ]
            + [
                f"2026-09-02 03:{m:02d}:00 INFO batch {m} complete latency_ms={90 + m}"
                for m in range(22, 30)
            ]
        ),
        encoding="utf-8",
    )
    put("engineering/ingest_service.py").write_text(
        '"""Ingest service."""\n\nDEFAULT_BATCH_TIMEOUT_S = 180\nMAX_RETRIES = 5\n\n\n'
        "def flush(batch):\n    return batch\n",
        encoding="utf-8",
    )
    put("engineering/meridian_release_notes.txt").write_text(
        "Meridian 4.3.0 release notes\n\nFixed ORR-1702 (webhook retries) and improved start-up time.\n",
        encoding="utf-8",
    )
    _image(
        put("specs/backplane_pinout.png"),
        IMAGE_TRANSCRIPTS["specs/backplane_pinout.png"],
        "PNG",
    )
    _image(put("support/asset_label.jpg"), IMAGE_TRANSCRIPTS["support/asset_label.jpg"], "JPEG")

    put("vendor/internal_audit.txt").write_text(
        "Internal audit notes: supplier terms were reviewed. The pricing schedule is held in the signed "
        "agreement and is not reproduced here.\n",
        encoding="utf-8",
    )
    put("vendor/supplier_agreement_ENCRYPTED.pdf").write_bytes(
        _encrypt_pdf(
            make_pdf(
                [["Supplier agreement", "Unit price of the TQ-40 at 10,000 units: USD 412.50"]]
            ),
            "s3cret",
        )
    )
    put("vendor/telemetry_capture.dat").write_bytes(os.urandom(2048))

    if include_unreadable:
        restricted = put("support/restricted_notes.txt")
        restricted.write_text(
            "Unit price of the TQ-40 at 10,000 units is 399.\n", encoding="utf-8"
        )
        restricted.chmod(0o000)
    return root


def image_key(path: Path) -> str:
    """Stable identity of an image file's bytes, used by the scripted vision model."""
    return hashlib.sha1(path.read_bytes(), usedforsecurity=False).hexdigest()


@dataclass
class ScriptedModel:
    """A stand-in for the model server: scripted answers and scripted image transcripts.

    ``replies`` maps a substring of the question to the JSON reply the "model" gives.
    ``transcripts`` maps the sha1 of an image's bytes to its transcription. Every call is
    recorded so tests can assert what the pipeline actually sent.
    """

    replies: dict[str, dict[str, Any]] = field(default_factory=dict)
    transcripts: dict[str, str] = field(default_factory=dict)
    ready: bool = True
    prompts: list[str] = field(default_factory=list)
    transcribe_calls: int = 0

    async def chat_json(
        self,
        system: str,
        user: str,
        schema: dict[str, Any],
        *,
        max_tokens: int,
        timeout_s: float,
    ) -> dict[str, Any] | None:
        """Return the scripted reply whose key occurs in the question."""
        self.prompts.append(user)
        question = user.split("\n", 1)[0]
        for key, reply in self.replies.items():
            if key in question:
                return reply
        return {"reasoning": "not found", "answerable": False, "answer": "", "sources": []}

    async def transcribe(self, image_bytes: bytes, *, timeout_s: float) -> str | None:
        """Return the scripted transcript for these image bytes."""
        self.transcribe_calls += 1
        return self.transcripts.get(
            hashlib.sha1(image_bytes, usedforsecurity=False).hexdigest()
        )

    async def wait_ready(self, budget_s: float) -> bool:
        """Report the scripted readiness."""
        return self.ready

    async def aclose(self) -> None:
        """Nothing to release."""


def good_replies() -> dict[str, dict[str, Any]]:
    """What a competent model says for each sample question, citing what it used."""

    def reply(answer: str, *sources: str) -> dict[str, Any]:
        return {
            "reasoning": "read the excerpt",
            "answerable": bool(answer),
            "answer": answer,
            "sources": list(sources),
        }

    return {
        "maximum junction temperature": reply("94 C", "specs/tq40_datasheet_r2.pdf"),
        "customer sampling": reply("Q3 FY27", "planning/roadmap_fy27.docx"),
        "fan assembly": reply("ORR-FAN-2214-B", "support/rma_parts.xlsx"),
        "fixed ticket ORR-1847": reply("4.3.2", "support/bug_database.csv"),
        "error code is logged": reply("E7731", "logs/prod_inference_2026-09-02.log"),
        "default batch timeout": reply("180", "engineering/ingest_service.py"),
        "THERM_ALERT#": reply("B14", "specs/backplane_pinout.png"),
        "board revision": reply("REV-C2", "support/asset_label.jpg"),
        "production log shows": reply(
            "4.3.2", "logs/prod_inference_2026-09-02.log", "support/bug_database.csv"
        ),
        "unit price": reply(""),
    }
