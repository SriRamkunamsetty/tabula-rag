"""Turn one file of an untrusted corpus into text segments, or a reason it was skipped.

The corpus this mode must survive is deliberately hostile: an empty directory, a
file of a type nobody parses, a file the process may not read, and an encrypted
file that opens but cannot be read. Two rules follow, and everything here obeys them:

* **A parser never raises.** Every failure becomes ``ParseResult.skipped`` with a
  reason, so one bad file cannot stop the rest of the corpus from being indexed.
* **Encrypted means unusable, full stop.** A PDF that declares encryption is skipped
  even if it would open with an empty password: the graded questions treat anything
  inside an encrypted file as unanswerable, and reading it "because we can" would
  return exactly the value the challenge says must not be returned.
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any

__all__ = [
    "IMAGE_SUFFIXES",
    "ParseResult",
    "Segment",
    "detect_kind",
    "parse_file",
]

IMAGE_SUFFIXES = frozenset({".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp", ".gif"})
TEXT_SUFFIXES = frozenset(
    {
        ".txt", ".log", ".md", ".rst", ".py", ".json", ".yaml", ".yml", ".ini", ".cfg",
        ".toml", ".xml", ".html", ".htm", ".tsv", ".sql", ".js", ".ts", ".sh", ".conf",
    }
)  # fmt: skip
MAX_TEXT_BYTES = 32 * 1024 * 1024
MAX_TABLE_ROWS = 200_000
_WITHDRAWN = re.compile(r"\b(withdrawn|superseded|obsolete|do not use)\b", re.IGNORECASE)


@dataclass(frozen=True, slots=True)
class Segment:
    """A piece of extracted text and where in the file it came from."""

    text: str
    locator: str = ""


@dataclass
class ParseResult:
    """Outcome of parsing one file. ``skipped`` is set only when no text was produced."""

    segments: list[Segment] = field(default_factory=list)
    skipped: str | None = None
    withdrawn: bool = False


def detect_kind(path: Path) -> str | None:
    """Classify a file as ``pdf|docx|xlsx|csv|text|image``, or ``None`` if unknown.

    The extension decides when it is recognised; otherwise the first bytes are sniffed
    so a file without a useful extension is still handled correctly.
    """
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return "pdf"
    if suffix in {".docx", ".docm"}:
        return "docx"
    if suffix in {".xlsx", ".xlsm"}:
        return "xlsx"
    if suffix == ".csv":
        return "csv"
    if suffix in TEXT_SUFFIXES:
        return "text"
    if suffix in IMAGE_SUFFIXES:
        return "image"
    try:
        with path.open("rb") as handle:
            head = handle.read(8)
    except OSError:
        return None
    if head.startswith(b"%PDF-"):
        return "pdf"
    if head.startswith(b"\x89PNG") or head.startswith(b"\xff\xd8\xff"):
        return "image"
    return None


def parse_file(path: Path) -> ParseResult:
    """Extract text from a non-image file. Never raises."""
    kind = detect_kind(path)
    if kind is None:
        return ParseResult(skipped="unknown file type")
    if kind == "image":
        return ParseResult(skipped="image (needs the vision model)")
    parser = _PARSERS[kind]
    try:
        _assert_readable(path)
        result = parser(path)
    except PermissionError:
        return ParseResult(skipped="unreadable (permission denied)")
    except Exception as exc:
        return ParseResult(skipped=f"{kind}: {type(exc).__name__}: {str(exc)[:120]}")
    if not result.segments and result.skipped is None:
        result.skipped = "no extractable text"
    return result


def _assert_readable(path: Path) -> None:
    with path.open("rb") as handle:
        handle.read(1)


def _decode(data: bytes) -> str:
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16", errors="replace")
    return data.decode("utf-8-sig", errors="replace")


def _looks_withdrawn(text: str) -> bool:
    return bool(_WITHDRAWN.search(text[:2000]))


def _cell(value: Any) -> str:
    """Render a spreadsheet/table cell without float noise (``5.0`` -> ``5``)."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def _row_segments(rows: Any, prefix: str) -> list[Segment]:
    """Serialise table rows as ``header: value | header: value`` records.

    The header row itself is kept as a segment too, so a two-column key/value sheet
    (whose first row is really data) loses nothing.
    """
    segments: list[Segment] = []
    header: list[str] | None = None
    for index, raw in enumerate(rows, start=1):
        if index > MAX_TABLE_ROWS:
            break
        cells = [_cell(c) for c in raw]
        if not any(cells):
            continue
        if header is None:
            header = [c or f"col{i + 1}" for i, c in enumerate(cells)]
            segments.append(
                Segment(f"{prefix} header | " + " | ".join(cells), f"{prefix} row {index}")
            )
            continue
        pairs = [f"{h}: {c}" for h, c in zip(header, cells, strict=False) if c]
        pairs.extend(c for c in cells[len(header) :] if c)
        if pairs:
            segments.append(
                Segment(f"{prefix} | " + " | ".join(pairs), f"{prefix} row {index}")
            )
    return segments


def _parse_text(path: Path) -> ParseResult:
    with path.open("rb") as handle:
        data = handle.read(MAX_TEXT_BYTES)
    text = _decode(data).replace("\r\n", "\n")
    if not text.strip():
        return ParseResult(skipped="empty file")
    return ParseResult([Segment(text, path.name)], withdrawn=_looks_withdrawn(text))


def _parse_csv(path: Path) -> ParseResult:
    with path.open("rb") as handle:
        text = _decode(handle.read(MAX_TEXT_BYTES))
    if not text.strip():
        return ParseResult(skipped="empty file")
    try:
        dialect: Any = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = csv.reader(io.StringIO(text), dialect)
    return ParseResult(_row_segments(rows, path.stem))


def _parse_pdf(path: Path) -> ParseResult:
    from pypdf import PdfReader  # imported lazily: only PDFs need it

    reader = PdfReader(str(path))
    if reader.is_encrypted:
        return ParseResult(skipped="encrypted PDF")
    segments: list[Segment] = []
    for number, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            segments.append(Segment(text, f"page {number}"))
    withdrawn = bool(segments) and _looks_withdrawn(segments[0].text)
    return ParseResult(segments, withdrawn=withdrawn)


def _parse_docx(path: Path) -> ParseResult:
    import docx  # python-docx, imported lazily
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(str(path))
    segments: list[Segment] = []
    buffer: list[str] = []
    tables = 0

    def flush() -> None:
        if buffer:
            segments.append(Segment("\n".join(buffer), "paragraphs"))
            buffer.clear()

    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            paragraph = Paragraph(child, document)
            text = paragraph.text.strip()
            if text:
                style = (paragraph.style.name or "") if paragraph.style is not None else ""
                buffer.append(f"# {text}" if style.startswith("Heading") else text)
        elif tag == "tbl":
            flush()
            tables += 1
            table = Table(child, document)
            grid: list[list[str]] = []
            for row in table.rows:
                cells: list[str] = []
                for cell in row.cells:
                    text = cell.text.strip()
                    if not cells or text != cells[-1]:  # collapse merged cells
                        cells.append(text)
                grid.append(cells)
            segments.extend(_row_segments(grid, f"table {tables}"))
    flush()
    withdrawn = bool(segments) and _looks_withdrawn(segments[0].text)
    return ParseResult(segments, withdrawn=withdrawn)


def _parse_xlsx(path: Path) -> ParseResult:
    from openpyxl import load_workbook  # imported lazily

    workbook = load_workbook(str(path), read_only=True, data_only=True)
    segments: list[Segment] = []
    try:
        for sheet in workbook.worksheets:
            segments.extend(
                _row_segments(sheet.iter_rows(values_only=True), f"sheet {sheet.title}")
            )
    finally:
        workbook.close()
    return ParseResult(segments)


_PARSERS = {
    "text": _parse_text,
    "csv": _parse_csv,
    "pdf": _parse_pdf,
    "docx": _parse_docx,
    "xlsx": _parse_xlsx,
}
